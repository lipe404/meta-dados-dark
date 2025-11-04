import os
import json
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify
from werkzeug.utils import secure_filename
from werkzeug.exceptions import RequestEntityTooLarge
import mimetypes
from datetime import datetime

# Bibliotecas para metadados
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS
import mutagen
from mutagen.id3 import ID3, TIT2, TPE1, TALB, TDRC, TCON
import PyPDF2
import docx
import openpyxl
import csv
import zipfile
import json as json_lib
import piexif
from fractions import Fraction

app = Flask(__name__)
app.config['SECRET_KEY'] = 'your-secret-key-here'
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max file size

# Criar pasta de uploads se não existir
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Tipos de arquivo permitidos
ALLOWED_EXTENSIONS = {
    'image': {'png', 'jpg', 'jpeg', 'gif', 'bmp', 'tiff', 'webp'},
    'audio': {'mp3', 'wav', 'flac', 'ogg', 'm4a', 'aac'},
    'video': {'mp4', 'avi', 'mkv', 'mov', 'wmv', 'flv', 'webm'},
    'document': {'pdf', 'docx', 'doc', 'xlsx', 'xls', 'pptx', 'txt'},
    'archive': {'zip', 'rar', '7z', 'tar', 'gz'}
}


def allowed_file(filename):
    """Verifica se o arquivo é permitido"""
    if '.' not in filename:
        return False
    ext = filename.rsplit('.', 1)[1].lower()
    for category, extensions in ALLOWED_EXTENSIONS.items():
        if ext in extensions:
            return True
    return False


def get_file_type(filename):
    """Determina o tipo do arquivo"""
    if '.' not in filename:
        return 'unknown'
    ext = filename.rsplit('.', 1)[1].lower()
    for category, extensions in ALLOWED_EXTENSIONS.items():
        if ext in extensions:
            return category
    return 'unknown'


class MetadataExtractor:
    """Classe para extrair metadados de diferentes tipos de arquivo"""

    @staticmethod
    def extract_image_metadata(filepath):
        """Extrai metadados de imagens com EXIF detalhado"""
        try:
            with Image.open(filepath) as image:
                metadata = {
                    'filename': os.path.basename(filepath),
                    'format': image.format,
                    'mode': image.mode,
                    'size': f"{image.size[0]}x{image.size[1]}",
                    'width': image.size[0],
                    'height': image.size[1]
                }

                # Extrair EXIF usando piexif para melhor suporte
                try:
                    exif_dict = piexif.load(filepath)
                except Exception:
                    exif_dict = {}

                # Informações básicas da imagem
                exif_info = {}

                # Mapeamento de tags EXIF mais relevantes
                if '0th' in exif_dict:
                    if piexif.ImageIFD.Make in exif_dict['0th']:
                        exif_info['camera_make'] = exif_dict['0th'][piexif.ImageIFD.Make].decode(
                            'utf-8', errors='ignore')
                    if piexif.ImageIFD.Model in exif_dict['0th']:
                        exif_info['camera_model'] = exif_dict['0th'][piexif.ImageIFD.Model].decode(
                            'utf-8', errors='ignore')
                    if piexif.ImageIFD.Software in exif_dict['0th']:
                        exif_info['software'] = exif_dict['0th'][piexif.ImageIFD.Software].decode(
                            'utf-8', errors='ignore')
                    if piexif.ImageIFD.Artist in exif_dict['0th']:
                        exif_info['artist'] = exif_dict['0th'][piexif.ImageIFD.Artist].decode(
                            'utf-8', errors='ignore')
                    if piexif.ImageIFD.Copyright in exif_dict['0th']:
                        exif_info['copyright'] = exif_dict['0th'][piexif.ImageIFD.Copyright].decode(
                            'utf-8', errors='ignore')

                # Configurações da câmera
                if 'Exif' in exif_dict:
                    if piexif.ExifIFD.DateTimeOriginal in exif_dict['Exif']:
                        exif_info['date_taken'] = exif_dict['Exif'][piexif.ExifIFD.DateTimeOriginal].decode(
                            'utf-8', errors='ignore')
                    if piexif.ExifIFD.ExposureTime in exif_dict['Exif']:
                        exp_time = exif_dict['Exif'][piexif.ExifIFD.ExposureTime]
                        if isinstance(exp_time, tuple):
                            exif_info['exposure_time'] = f"1/{int(exp_time[1]/exp_time[0])}s" if exp_time[
                                1] > exp_time[0] else f"{exp_time[0]/exp_time[1]}s"
                        else:
                            exif_info['exposure_time'] = f"{exp_time}s"
                    if piexif.ExifIFD.FNumber in exif_dict['Exif']:
                        fnum = exif_dict['Exif'][piexif.ExifIFD.FNumber]
                        if isinstance(fnum, tuple):
                            exif_info['f_number'] = f"f/{fnum[0]/fnum[1]}"
                        else:
                            exif_info['f_number'] = f"f/{fnum}"
                    if piexif.ExifIFD.ISOSpeedRatings in exif_dict['Exif']:
                        exif_info['iso'] = str(
                            exif_dict['Exif'][piexif.ExifIFD.ISOSpeedRatings])
                    if piexif.ExifIFD.FocalLength in exif_dict['Exif']:
                        focal = exif_dict['Exif'][piexif.ExifIFD.FocalLength]
                        if isinstance(focal, tuple):
                            exif_info['focal_length'] = f"{focal[0]/focal[1]}mm"
                        else:
                            exif_info['focal_length'] = f"{focal}mm"
                    if piexif.ExifIFD.Flash in exif_dict['Exif']:
                        exif_info['flash'] = "Sim" if exif_dict['Exif'][piexif.ExifIFD.Flash] & 0x1 else "Não"
                    if piexif.ExifIFD.WhiteBalance in exif_dict['Exif']:
                        exif_info['white_balance'] = "Automático" if exif_dict['Exif'][piexif.ExifIFD.WhiteBalance] == 0 else "Manual"

                # Informações da lente
                if 'Exif' in exif_dict:
                    if piexif.ExifIFD.LensMake in exif_dict['Exif']:
                        exif_info['lens_make'] = exif_dict['Exif'][piexif.ExifIFD.LensMake].decode(
                            'utf-8', errors='ignore')
                    if piexif.ExifIFD.LensModel in exif_dict['Exif']:
                        exif_info['lens_model'] = exif_dict['Exif'][piexif.ExifIFD.LensModel].decode(
                            'utf-8', errors='ignore')

                # Informações de localização GPS
                if 'GPS' in exif_dict:
                    gps_info = {}
                    if piexif.GPSIFD.GPSLatitude in exif_dict['GPS'] and piexif.GPSIFD.GPSLatitudeRef in exif_dict['GPS']:
                        lat = exif_dict['GPS'][piexif.GPSIFD.GPSLatitude]
                        lat_ref = exif_dict['GPS'][piexif.GPSIFD.GPSLatitudeRef].decode(
                            'utf-8')
                        if isinstance(lat[0], tuple):
                            lat_deg = lat[0][0]/lat[0][1] + lat[1][0] / \
                                lat[1][1]/60 + lat[2][0]/lat[2][1]/3600
                            gps_info['latitude'] = f"{lat_deg:.6f}° {lat_ref}"
                        else:
                            gps_info['latitude'] = f"{lat[0]}° {lat_ref}"

                    if piexif.GPSIFD.GPSLongitude in exif_dict['GPS'] and piexif.GPSIFD.GPSLongitudeRef in exif_dict['GPS']:
                        lon = exif_dict['GPS'][piexif.GPSIFD.GPSLongitude]
                        lon_ref = exif_dict['GPS'][piexif.GPSIFD.GPSLongitudeRef].decode(
                            'utf-8')
                        if isinstance(lon[0], tuple):
                            lon_deg = lon[0][0]/lon[0][1] + lon[1][0] / \
                                lon[1][1]/60 + lon[2][0]/lon[2][1]/3600
                            gps_info['longitude'] = f"{lon_deg:.6f}° {lon_ref}"
                        else:
                            gps_info['longitude'] = f"{lon[0]}° {lon_ref}"

                    if gps_info:
                        exif_info['gps'] = gps_info

                # Todos os dados EXIF brutos (para metadados ocultos)
                if exif_dict:
                    metadata['exif_raw'] = {}
                    for ifd in exif_dict:
                        if ifd != 'thumbnail':
                            metadata['exif_raw'][ifd] = {}
                            for tag in exif_dict[ifd]:
                                tag_name = TAGS.get(
                                    tag, tag) if ifd == '0th' or ifd == '1st' else tag
                                value = exif_dict[ifd][tag]
                                if isinstance(value, bytes):
                                    try:
                                        value = value.decode(
                                            'utf-8', errors='ignore')
                                    except:
                                        value = str(value)
                                elif isinstance(value, tuple) and len(value) > 0:
                                    if isinstance(value[0], tuple):
                                        # Fração
                                        try:
                                            value = str(
                                                value[0][0] / value[0][1])
                                        except:
                                            value = str(value)
                                    else:
                                        value = str(value)
                                else:
                                    value = str(value)
                                metadata['exif_raw'][ifd][str(
                                    tag_name)] = value

                if exif_info:
                    metadata['exif'] = exif_info

                return metadata
        except Exception as e:
            return {'error': str(e)}

    @staticmethod
    def extract_audio_metadata(filepath):
        """Extrai metadados de áudio"""
        try:
            audiofile = mutagen.File(filepath)
            if audiofile is None:
                return {'error': 'Arquivo de áudio não suportado'}

            metadata = {
                'filename': os.path.basename(filepath),
                'length': f"{audiofile.info.length:.2f} segundos" if hasattr(audiofile.info, 'length') else 'N/A',
                'bitrate': f"{audiofile.info.bitrate} bps" if hasattr(audiofile.info, 'bitrate') else 'N/A',
                'sample_rate': f"{audiofile.info.sample_rate} Hz" if hasattr(audiofile.info, 'sample_rate') else 'N/A',
                'channels': audiofile.info.channels if hasattr(audiofile.info, 'channels') else 'N/A'
            }

            # Tags comuns
            common_tags = {
                'title': ['TIT2', 'TITLE', '\xa9nam'],
                'artist': ['TPE1', 'ARTIST', '\xa9ART'],
                'album': ['TALB', 'ALBUM', '\xa9alb'],
                'date': ['TDRC', 'DATE', '\xa9day'],
                'genre': ['TCON', 'GENRE', '\xa9gen']
            }

            for tag_name, possible_keys in common_tags.items():
                for key in possible_keys:
                    if key in audiofile:
                        value = audiofile[key]
                        if isinstance(value, list) and value:
                            value = str(value[0])
                        metadata[tag_name] = str(value)
                        break
                else:
                    metadata[tag_name] = 'N/A'

            return metadata
        except Exception as e:
            return {'error': str(e)}

    @staticmethod
    def extract_pdf_metadata(filepath):
        """Extrai metadados de PDFs"""
        try:
            with open(filepath, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                metadata = {
                    'filename': os.path.basename(filepath),
                    'pages': len(pdf_reader.pages),
                    'encrypted': pdf_reader.is_encrypted
                }

                if pdf_reader.metadata:
                    pdf_metadata = pdf_reader.metadata
                    metadata.update({
                        'title': pdf_metadata.get('/Title', 'N/A'),
                        'author': pdf_metadata.get('/Author', 'N/A'),
                        'subject': pdf_metadata.get('/Subject', 'N/A'),
                        'creator': pdf_metadata.get('/Creator', 'N/A'),
                        'producer': pdf_metadata.get('/Producer', 'N/A'),
                        'creation_date': str(pdf_metadata.get('/CreationDate', 'N/A')),
                        'modification_date': str(pdf_metadata.get('/ModDate', 'N/A'))
                    })

                return metadata
        except Exception as e:
            return {'error': str(e)}

    @staticmethod
    def extract_document_metadata(filepath):
        """Extrai metadados de documentos Office"""
        try:
            ext = filepath.rsplit('.', 1)[1].lower()
            metadata = {'filename': os.path.basename(filepath)}

            if ext == 'docx':
                doc = docx.Document(filepath)
                core_props = doc.core_properties
                metadata.update({
                    'title': core_props.title or 'N/A',
                    'author': core_props.author or 'N/A',
                    'subject': core_props.subject or 'N/A',
                    'created': str(core_props.created) if core_props.created else 'N/A',
                    'modified': str(core_props.modified) if core_props.modified else 'N/A',
                    'paragraphs': len(doc.paragraphs)
                })

            elif ext in ['xlsx', 'xls']:
                workbook = openpyxl.load_workbook(filepath)
                props = workbook.properties
                metadata.update({
                    'title': props.title or 'N/A',
                    'creator': props.creator or 'N/A',
                    'subject': props.subject or 'N/A',
                    'created': str(props.created) if props.created else 'N/A',
                    'modified': str(props.modified) if props.modified else 'N/A',
                    'worksheets': len(workbook.worksheets)
                })

            return metadata
        except Exception as e:
            return {'error': str(e)}

    @staticmethod
    def extract_general_metadata(filepath):
        """Extrai metadados gerais do arquivo"""
        try:
            stat = os.stat(filepath)
            metadata = {
                'filename': os.path.basename(filepath),
                'size': f"{stat.st_size / 1024:.2f} KB",
                'created': datetime.fromtimestamp(stat.st_ctime).strftime('%Y-%m-%d %H:%M:%S'),
                'modified': datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M:%S'),
                'accessed': datetime.fromtimestamp(stat.st_atime).strftime('%Y-%m-%d %H:%M:%S'),
                'mime_type': mimetypes.guess_type(filepath)[0] or 'unknown'
            }
            return metadata
        except Exception as e:
            return {'error': str(e)}


def extract_metadata(filepath):
    """Extrai metadados baseado no tipo do arquivo"""
    file_type = get_file_type(filepath)

    if file_type == 'image':
        return MetadataExtractor.extract_image_metadata(filepath)
    elif file_type == 'audio':
        return MetadataExtractor.extract_audio_metadata(filepath)
    elif file_type == 'document':
        if filepath.lower().endswith('.pdf'):
            return MetadataExtractor.extract_pdf_metadata(filepath)
        else:
            return MetadataExtractor.extract_document_metadata(filepath)
    else:
        return MetadataExtractor.extract_general_metadata(filepath)


@app.route('/')
def index():
    """Página inicial"""
    uploaded_files = []
    upload_dir = app.config['UPLOAD_FOLDER']
    if os.path.exists(upload_dir):
        for filename in os.listdir(upload_dir):
            if allowed_file(filename):
                filepath = os.path.join(upload_dir, filename)
                file_info = {
                    'name': filename,
                    'type': get_file_type(filename),
                    'size': f"{os.path.getsize(filepath) / 1024:.2f} KB"
                }
                uploaded_files.append(file_info)

    return render_template('index.html', files=uploaded_files)


@app.route('/upload', methods=['POST'])
def upload_file():
    """Rota para upload de arquivos"""
    try:
        if 'file' not in request.files:
            flash('Nenhum arquivo selecionado', 'error')
            return redirect(url_for('index'))

        file = request.files['file']
        if file.filename == '':
            flash('Nenhum arquivo selecionado', 'error')
            return redirect(url_for('index'))

        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            # Evitar sobrescrever arquivos
            counter = 1
            base_name, ext = os.path.splitext(filename)
            while os.path.exists(os.path.join(app.config['UPLOAD_FOLDER'], filename)):
                filename = f"{base_name}_{counter}{ext}"
                counter += 1

            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(filepath)
            flash(f'Arquivo {filename} enviado com sucesso!', 'success')
            return redirect(url_for('view_metadata', filename=filename))
        else:
            flash('Tipo de arquivo não permitido', 'error')
            return redirect(url_for('index'))

    except RequestEntityTooLarge:
        flash('Arquivo muito grande. Tamanho máximo: 50MB', 'error')
        return redirect(url_for('index'))
    except Exception as e:
        flash(f'Erro no upload: {str(e)}', 'error')
        return redirect(url_for('index'))


@app.route('/metadata/<filename>')
def view_metadata(filename):
    """Visualizar metadados do arquivo"""
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    if not os.path.exists(filepath):
        flash('Arquivo não encontrado', 'error')
        return redirect(url_for('index'))

    metadata = extract_metadata(filepath)
    file_type = get_file_type(filename)

    return render_template('metadata.html',
                           filename=filename,
                           metadata=metadata,
                           file_type=file_type)


@app.route('/edit_metadata/<filename>')
def edit_metadata(filename):
    """Página para editar metadados"""
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    if not os.path.exists(filepath):
        flash('Arquivo não encontrado', 'error')
        return redirect(url_for('index'))

    metadata = extract_metadata(filepath)
    file_type = get_file_type(filename)

    return render_template('edit_metadata.html',
                           filename=filename,
                           metadata=metadata,
                           file_type=file_type)


@app.route('/save_metadata/<filename>', methods=['POST'])
def save_metadata(filename):
    """Salvar metadados editados"""
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    if not os.path.exists(filepath):
        flash('Arquivo não encontrado', 'error')
        return redirect(url_for('index'))

    try:
        file_type = get_file_type(filename)

        if file_type == 'audio':
            # Editar metadados de áudio
            audiofile = mutagen.File(filepath)
            if audiofile is not None:
                # Limpar tags existentes
                audiofile.clear()

                # Adicionar novas tags
                if 'title' in request.form and request.form['title']:
                    audiofile['TIT2'] = TIT2(
                        encoding=3, text=request.form['title'])
                if 'artist' in request.form and request.form['artist']:
                    audiofile['TPE1'] = TPE1(
                        encoding=3, text=request.form['artist'])
                if 'album' in request.form and request.form['album']:
                    audiofile['TALB'] = TALB(
                        encoding=3, text=request.form['album'])
                if 'date' in request.form and request.form['date']:
                    audiofile['TDRC'] = TDRC(
                        encoding=3, text=request.form['date'])
                if 'genre' in request.form and request.form['genre']:
                    audiofile['TCON'] = TCON(
                        encoding=3, text=request.form['genre'])

                audiofile.save()
                flash('Metadados de áudio salvos com sucesso!', 'success')

        elif file_type == 'image':
            # Editar metadados EXIF de imagens
            try:
                exif_dict = piexif.load(filepath)
                if not exif_dict:
                    exif_dict = {'0th': {}, 'Exif': {}, 'GPS': {}, '1st': {}}

                # Editar informações básicas
                if 'artist' in request.form and request.form['artist']:
                    exif_dict['0th'][piexif.ImageIFD.Artist] = request.form['artist'].encode(
                        'utf-8')
                if 'copyright' in request.form and request.form['copyright']:
                    exif_dict['0th'][piexif.ImageIFD.Copyright] = request.form['copyright'].encode(
                        'utf-8')
                if 'software' in request.form and request.form['software']:
                    exif_dict['0th'][piexif.ImageIFD.Software] = request.form['software'].encode(
                        'utf-8')
                if 'camera_make' in request.form and request.form['camera_make']:
                    exif_dict['0th'][piexif.ImageIFD.Make] = request.form['camera_make'].encode(
                        'utf-8')
                if 'camera_model' in request.form and request.form['camera_model']:
                    exif_dict['0th'][piexif.ImageIFD.Model] = request.form['camera_model'].encode(
                        'utf-8')

                # Editar configurações da câmera
                if 'date_taken' in request.form and request.form['date_taken']:
                    date_str = request.form['date_taken'].replace(
                        ' ', ':').replace('-', ':')
                    exif_dict['Exif'][piexif.ExifIFD.DateTimeOriginal] = date_str.encode(
                        'utf-8')
                    exif_dict['Exif'][piexif.ExifIFD.DateTimeDigitized] = date_str.encode(
                        'utf-8')
                    exif_dict['0th'][piexif.ImageIFD.DateTime] = date_str.encode(
                        'utf-8')

                # Editar informações da lente
                if 'lens_make' in request.form and request.form['lens_make']:
                    exif_dict['Exif'][piexif.ExifIFD.LensMake] = request.form['lens_make'].encode(
                        'utf-8')
                if 'lens_model' in request.form and request.form['lens_model']:
                    exif_dict['Exif'][piexif.ExifIFD.LensModel] = request.form['lens_model'].encode(
                        'utf-8')

                # Editar GPS (latitude e longitude)
                if 'gps_latitude' in request.form and request.form['gps_latitude']:
                    try:
                        lat = float(request.form['gps_latitude'])
                        lat_ref = 'N' if lat >= 0 else 'S'
                        lat = abs(lat)
                        lat_deg = int(lat)
                        lat_min = int((lat - lat_deg) * 60)
                        lat_sec = ((lat - lat_deg) * 60 - lat_min) * 60
                        exif_dict['GPS'][piexif.GPSIFD.GPSLatitude] = (
                            (int(lat_deg * 100), 100),
                            (int(lat_min * 100), 100),
                            (int(lat_sec * 10000), 10000)
                        )
                        exif_dict['GPS'][piexif.GPSIFD.GPSLatitudeRef] = lat_ref.encode(
                            'utf-8')
                    except ValueError:
                        pass

                if 'gps_longitude' in request.form and request.form['gps_longitude']:
                    try:
                        lon = float(request.form['gps_longitude'])
                        lon_ref = 'E' if lon >= 0 else 'W'
                        lon = abs(lon)
                        lon_deg = int(lon)
                        lon_min = int((lon - lon_deg) * 60)
                        lon_sec = ((lon - lon_deg) * 60 - lon_min) * 60
                        exif_dict['GPS'][piexif.GPSIFD.GPSLongitude] = (
                            (int(lon_deg * 100), 100),
                            (int(lon_min * 100), 100),
                            (int(lon_sec * 10000), 10000)
                        )
                        exif_dict['GPS'][piexif.GPSIFD.GPSLongitudeRef] = lon_ref.encode(
                            'utf-8')
                    except ValueError:
                        pass

                # Salvar EXIF de volta no arquivo
                exif_bytes = piexif.dump(exif_dict)
                piexif.insert(exif_bytes, filepath)

                flash('Metadados EXIF salvos com sucesso!', 'success')
            except Exception as e:
                flash(f'Erro ao salvar metadados EXIF: {str(e)}', 'error')

        else:
            flash(
                'Edição de metadados não suportada para este tipo de arquivo', 'warning')

        return redirect(url_for('view_metadata', filename=filename))

    except Exception as e:
        flash(f'Erro ao salvar metadados: {str(e)}', 'error')
        return redirect(url_for('edit_metadata', filename=filename))


@app.route('/delete_file/<filename>', methods=['POST'])
def delete_file(filename):
    """Deletar arquivo"""
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    try:
        if os.path.exists(filepath):
            os.remove(filepath)
            flash(f'Arquivo {filename} deletado com sucesso!', 'success')
        else:
            flash('Arquivo não encontrado', 'error')
    except Exception as e:
        flash(f'Erro ao deletar arquivo: {str(e)}', 'error')

    return redirect(url_for('index'))


@app.errorhandler(413)
def too_large(e):
    flash('Arquivo muito grande. Tamanho máximo: 50MB', 'error')
    return redirect(url_for('index'))


if __name__ == '__main__':
    app.run(debug=True)
